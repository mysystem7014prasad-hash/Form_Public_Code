from django.shortcuts import render
from django.conf import settings
from .forms import DetailsForm

from openpyxl import Workbook, load_workbook
import os


def details_form(request):

    if request.method == "POST":

        form = DetailsForm(request.POST)

        if form.is_valid():

            rollno = form.cleaned_data["rollno"]
            name = form.cleaned_data["name"]
            email = form.cleaned_data["email"]
            phone = form.cleaned_data["phone"]
            address = form.cleaned_data["address"]
            course = form.cleaned_data["course"]

            # Excel file location
            file_path = os.path.join(settings.BASE_DIR, "details.xlsx")

            # Create Excel file if it doesn't exist
            if not os.path.exists(file_path):

                workbook = Workbook()
                sheet = workbook.active

                sheet.append([
                    "Roll No",
                    "Name",
                    "Email",
                    "Phone",
                    "Address",
                    "Course"
                ])

                workbook.save(file_path)

            # Open Excel file
            workbook = load_workbook(file_path)
            sheet = workbook.active

            # Check whether Roll No already exists
            rollno_exists = False

            for row in sheet.iter_rows(min_row=2, values_only=True):
                if str(row[0]) == str(rollno):
                    rollno_exists = True
                    break

            # If Roll No already exists
            if rollno_exists:

                form.add_error(
                    "rollno",
                    "This Roll No is already registered."
                )

                return render(
                    request,
                    "details/details_form.html",
                    {"form": form}
                )

            # Add new student
            sheet.append([
                rollno,
                name,
                email,
                phone,
                address,
                course
            ])

            # Save Excel
            workbook.save(file_path)

            # Show success page
            return render(
                request,
                "details/success.html"
            )

    else:
        form = DetailsForm()

    return render(
        request,
        "details/details_form.html",
        {"form": form}
    )